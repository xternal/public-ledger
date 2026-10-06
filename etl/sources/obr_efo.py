"""
OBR Economic and fiscal outlook (EFO), latest edition.

Discovery (fetch): https://obr.uk/efo/ redirects to the newest EFO page. Its
canonical URL and every "/efo/economic-and-fiscal-outlook-<month>-<year>/" link
on it are candidate editions; the newest (by year, month) whose page links both
workbooks below is used. Title, published date and vintage ("EFO-YYYY-MM") are
read from that page. If discovery fails, the March 2026 EFO is used and a
warning is logged; the edition used is always logged (logger "etl.sources.obr_efo")
and kept in EDITION, and SOURCE is updated in place to describe it.

Workbooks used (found on the edition page by link pattern):

  * "Charts and tables: annex tables"
        Table A.3  Determinants of the fiscal forecast  -> nominal GDP (fiscal year, and centred end-March)
        Table A.5  Current receipts                     -> receipts.total, the 13 Statement lines, receipts.detail.*
        Table A.7  Total managed expenditure            -> spending.tme, spending.debt_interest
        Table A.9  Fiscal aggregates                    -> fiscal.psnb, fiscal.psnd, fiscal.psnd_pct_gdp
  * "Detailed forecast tables: debt interest ready reckoner"
        Table 5.1  Debt interest (net of APF) ready reckoner -> reckoner.debt_interest.*

Sheets are found by name (TA.5 ...) and checked by their title; values are read
by row label and by the fiscal-year header of their column, never by cell
position. A missing sheet, row or year raises. Columns at or after the
"Forecast" marker are kind "forecast"; earlier ones are "outturn".
"""

from __future__ import annotations

import calendar
import json
import logging
import re
from dataclasses import dataclass
from datetime import date
from html import unescape
from urllib.parse import urljoin, urlparse, urlunparse

import httpx
import openpyxl

from etl.core import (
    RAW_DIR,
    Observation,
    RawArtifact,
    Source,
    download,
    links,
    normalise_fiscal_year,
)

log = logging.getLogger(__name__)

INDEX_URL = "https://obr.uk/efo/"  # redirects to the newest EFO
FALLBACK_URL = "https://obr.uk/efo/economic-and-fiscal-outlook-march-2026/"
MAX_EDITION_PAGES = 2  # politeness: at most this many edition pages tried per run

SOURCE = Source(
    id="obr_efo",
    title="OBR Economic and fiscal outlook – March 2026",
    publisher="Office for Budget Responsibility",
    url=FALLBACK_URL,
    licence="OGL v3",
    published_on=date(2026, 3, 3),
    # Two EFOs a year, but not evenly spaced: with the Budget in late November the
    # gap from a spring EFO to the autumn one is about eight months, so 183 days
    # would fail CI in mid-October for no real problem. fetch() sets published_on
    # from the discovered edition, so staleness counts from the real date.
    cadence_days=245,
    grace_days=30,
)

# role -> (regex for the download link on the edition page, local filename suffix)
FILES = {
    "annex": (r"/download/[^/?#]*charts-and-tables-annex-tables/", "annex_tables.xlsx"),
    "debt_interest_rr": (r"/download/[^/?#]*debt-interest-ready-reckoner[^/?#]*/", "debt_interest_ready_reckoner.xlsx"),
}

# Statement receipt lines -> exact (cleaned, lower-case) labels of Table A.5 top-level rows.
# Every remaining A.5 tax row goes to other_taxes; every non-tax row goes to non_tax.
RECEIPT_LINES: dict[str, list[str]] = {
    "income_tax": ["income tax"],
    "nics": ["national insurance contributions"],
    "vat": ["value added tax"],
    "corp_tax": ["corporation tax"],
    "council_tax": ["council tax"],
    "business_rates": ["business rates"],
    "fuel_duty": ["fuel duties"],
    "stamp_duty": ["property transaction taxes", "stamp taxes on shares"],
    "cgt": ["capital gains tax"],
    "alcohol_tobacco": ["alcohol duties", "tobacco duties"],
    "iht": ["inheritance tax"],
}
STATEMENT_LINES = [*RECEIPT_LINES, "other_taxes", "non_tax"]

NA_TAXES = "national accounts taxes"
CURRENT_RECEIPTS = "current receipts"

# Sums of published rows must match the published subtotal to this many £bn.
SUM_TOLERANCE_BN = 0.01


# --------------------------------------------------------------------------- discovery

_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
_EDITION_PATH = re.compile(r"^/efo/economic-and-fiscal-outlook-([a-z]+)-(\d{4})/?$")
_FILE_NAME = re.compile(r"^efo_(\d{4})_(\d{2})_(.+)$")


@dataclass
class Edition:
    url: str
    year: int
    month: int
    title: str
    published_on: date | None
    files: dict[str, str]  # role -> workbook download URL
    discovered: bool = True

    @property
    def vintage(self) -> str:
        return f"EFO-{self.year}-{self.month:02d}"

    @property
    def prefix(self) -> str:
        return f"efo_{self.year}_{self.month:02d}"

    @property
    def label(self) -> str:
        return f"{calendar.month_name[self.month]} {self.year}"


EDITION: Edition | None = None  # set by fetch()


def _strip_query(url: str) -> str:
    return urlunparse(urlparse(url)._replace(query="", fragment=""))


def edition_of(url: str) -> tuple[int, int] | None:
    """'https://obr.uk/efo/economic-and-fiscal-outlook-march-2026/' -> (2026, 3)."""
    u = urlparse(url)
    if u.netloc and u.netloc not in ("obr.uk", "www.obr.uk"):
        return None
    m = _EDITION_PATH.match(u.path)
    if not m or m.group(1) not in _MONTHS:
        return None
    return int(m.group(2)), _MONTHS[m.group(1)]


def _canonical(html: str, base: str) -> str | None:
    m = re.search(r'<link[^>]+rel="canonical"[^>]+href="([^"]+)"', html)
    return urljoin(base, unescape(m.group(1))) if m else None


def edition_candidates(html: str, base: str) -> list[str]:
    """Edition page URLs named on a page (canonical first), newest first, without duplicates."""
    found = {}
    for u in [_canonical(html, base) or "", *links(html, base, r"/efo/economic-and-fiscal-outlook-[a-z]+-\d{4}")]:
        u = _strip_query(u)
        if not u.endswith("/"):
            u += "/"
        ym = edition_of(u)
        if ym and u not in found:
            found[u] = ym
    return sorted(found, key=lambda u: found[u], reverse=True)


def page_title(html: str) -> str | None:
    m = re.search(r'<meta[^>]+property="og:title"[^>]+content="([^"]+)"', html) or re.search(r"<title>([^<]+)</title>", html)
    if not m:
        return None
    t = re.sub(r"\s+[-|–]\s+Office for Budget Responsibility\s*$", "", unescape(m.group(1))).strip()
    return f"OBR {t}"


def page_published_on(html: str) -> date | None:
    m = re.search(r'"datePublished"\s*:\s*"(\d{4}-\d{2}-\d{2})', html) or re.search(
        r'<meta[^>]+property="article:published_time"[^>]+content="(\d{4}-\d{2}-\d{2})', html
    )
    return date.fromisoformat(m.group(1)) if m else None


def edition_from_page(url: str, html: str, discovered: bool = True) -> Edition:
    """Read an edition page; raises LookupError if it does not link both workbooks."""
    ym = edition_of(url)
    if ym is None:
        raise LookupError(f"{url} is not an EFO edition page")
    year, month = ym
    tag = f"{calendar.month_name[month].lower()}-{year}"
    files = {}
    for role, (pattern, _) in FILES.items():
        found = [_strip_query(u) for u in links(html, url, pattern)]
        if not found:
            raise LookupError(f"{url}: no {role} workbook link matching {pattern!r}")
        # Prefer this edition's own file; parse() checks the workbook's edition anyway.
        own = [u for u in found if tag in u]
        files[role] = (own or found)[0]
    title = page_title(html) or f"OBR Economic and fiscal outlook – {calendar.month_name[month]} {year}"
    published = page_published_on(html)
    if published is None:
        log.warning("obr_efo: no publication date on %s", url)
    return Edition(url, year, month, title, published, files, discovered)


def _page(url: str, filename: str) -> str:
    return download(SOURCE.id, url, filename).path.read_text(errors="replace")


def discover() -> Edition:
    """The newest EFO edition that links both workbooks, else the March 2026 fallback."""
    try:
        index_html = _page(INDEX_URL, "efo_index.html")
        index_url = _canonical(index_html, INDEX_URL)
        candidates = edition_candidates(index_html, INDEX_URL)
        if not candidates:
            raise LookupError(f"no EFO edition links on {INDEX_URL}")
        for url in candidates[:MAX_EDITION_PAGES]:
            try:
                if url == index_url:
                    html = index_html  # /efo/ already served this edition's page
                else:
                    y, m = edition_of(url)
                    html = _page(url, f"efo_{y}_{m:02d}_landing.html")
                return edition_from_page(url, html)
            except (LookupError, httpx.HTTPError) as e:
                log.warning("obr_efo: skipping %s: %s", url, e)
        raise LookupError(f"none of {candidates[:MAX_EDITION_PAGES]} links both workbooks")
    except (LookupError, httpx.HTTPError, OSError) as e:
        log.warning("obr_efo: discovery failed (%s); falling back to %s", e, FALLBACK_URL)
    html = _page(FALLBACK_URL, "efo_2026_03_landing.html")
    return edition_from_page(FALLBACK_URL, html, discovered=False)


def fetch(since: date | None = None) -> list[RawArtifact]:
    """
    Discover the latest EFO, update SOURCE to describe it, and download its two
    workbooks. `since` is accepted for the common contract and not used: an EFO
    is one edition, and core.download() caches for 20 h.
    """
    global EDITION
    ed = discover()
    EDITION = ed
    SOURCE.url, SOURCE.title, SOURCE.published_on = ed.url, ed.title, ed.published_on
    log.info("obr_efo: using %s (%s, published %s, %s) from %s", ed.vintage, ed.title, ed.published_on,
             "discovered" if ed.discovered else "fallback", ed.url)
    return [
        download(SOURCE.id, ed.files[role], f"{ed.prefix}_{suffix}", vintage=ed.vintage)
        for role, (_, suffix) in FILES.items()
    ]


def _vintage_of(path) -> str:
    m = _FILE_NAME.match(path.name)
    if not m:
        raise ValueError(f"obr_efo: cannot tell the edition of {path.name}")
    return f"EFO-{m.group(1)}-{m.group(2)}"


def cached_artifacts() -> list[RawArtifact] | None:
    """The newest edition's workbooks already in data/raw/obr_efo/ (for offline runs), or None."""
    folder = RAW_DIR / SOURCE.id
    suffixes = [suffix for _, suffix in FILES.values()]
    prefixes = sorted({p.name[: -len(suffixes[0]) - 1] for p in folder.glob(f"efo_*_{suffixes[0]}")}, reverse=True)
    for prefix in prefixes:
        paths = [folder / f"{prefix}_{s}" for s in suffixes]
        metas = [p.parent / f"{p.name}.meta.json" for p in paths]
        if all(p.exists() for p in paths + metas):
            out = []
            for path, meta_path in zip(paths, metas):
                meta = json.loads(meta_path.read_text())
                out.append(RawArtifact(SOURCE.id, meta.get("requested_url") or meta["url"], path, meta["sha256"],
                                       meta["fetched_at"], meta.get("content_type"), _vintage_of(path)))
            return out
    return None


# --------------------------------------------------------------------------- table reading

_FOOTNOTE = re.compile(r"(?<=[A-Za-z)])\d+(?:,\d+)*$")


def clean_label(text) -> str:
    """'Income tax1' -> 'Income tax'; 'Nominal GDP (£ billion)1,2' -> 'Nominal GDP (£ billion)'."""
    s = re.sub(r"\s+", " ", str(text)).strip()
    return _FOOTNOTE.sub("", s).strip()


def slug(text: str) -> str:
    s = clean_label(text).lower().replace("&", " and ").replace("£", " gbp ")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


class Row:
    def __init__(self, index: int, labels: list[tuple[int, str]], values: dict[str, float], section: str | None):
        self.index = index
        self.labels = labels  # (column, cleaned text), left to right
        self.values = values  # period -> value
        self.section = section  # last unit marker seen above the row (e.g. "£ billion")

    @property
    def label(self) -> str:
        texts = [t for _, t in self.labels if not t.lower().startswith("of which")]
        return texts[-1] if texts else ""

    @property
    def key(self) -> str:
        return self.label.lower()


class Table:
    def __init__(self, ws, vintage: str):
        self.title = ws.title
        self.vintage = vintage
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        header = None
        for i, r in enumerate(grid):
            years = {j: normalise_fiscal_year(v) for j, v in enumerate(r) if isinstance(v, str) and normalise_fiscal_year(v)}
            if len(years) >= 3:
                header = i
                break
        if header is None:
            raise ValueError(f"{ws.title}: no fiscal-year header row")
        self.periods: dict[int, str] = years
        first_year_col = min(years)

        forecast_col = None
        for k in range(max(0, header - 3), header):
            for j, v in enumerate(grid[k]):
                if isinstance(v, str) and v.strip().lower().startswith("forecast"):
                    forecast_col = j
        if forecast_col is None:
            raise ValueError(f"{ws.title}: no 'Forecast' marker above the year header")
        self.kinds = {p: ("forecast" if j >= forecast_col else "outturn") for j, p in years.items()}

        self.rows: list[Row] = []
        section = None
        for i in range(len(grid)):
            r = grid[i]
            texts = [(j, clean_label(v)) for j, v in enumerate(r) if isinstance(v, str) and v.strip()]
            for _, t in texts:
                if t.lower() in ("£ billion", "per cent of gdp"):
                    section = t.lower()  # unit markers may sit above the header too
            if i <= header:
                continue
            labels = [(j, t) for j, t in texts if j < first_year_col]
            values = {}
            for j, p in years.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    values[p] = float(v)
            self.rows.append(Row(i + 1, labels, values, section))

    def data_rows(self) -> list[Row]:
        return [r for r in self.rows if r.values and r.label]

    def find(self, label: str, *, section: str | None = None, startswith: bool = False) -> Row:
        want = label.lower()
        hits = [
            r for r in self.data_rows()
            if (r.key.startswith(want) if startswith else r.key == want) and (section is None or r.section == section)
        ]
        if len(hits) != 1:
            raise ValueError(f"{self.title}: expected one row {label!r} (section {section}), found {len(hits)}")
        row = hits[0]
        missing = set(self.periods.values()) - set(row.values)
        if missing:
            raise ValueError(f"{self.title}: row {label!r} has no value for {sorted(missing)}")
        return row


def _edition_label(vintage: str) -> str:
    """'EFO-2026-03' -> 'March 2026'."""
    y, m = vintage.removeprefix("EFO-").split("-")
    return f"{calendar.month_name[int(m)]} {y}"


def _workbook(raw: RawArtifact, vintage: str):
    wb = openpyxl.load_workbook(raw.path, read_only=True, data_only=True)
    if "Contents" not in wb.sheetnames:
        raise ValueError(f"{raw.path.name}: no Contents sheet")
    contents = " ".join(str(v) for r in wb["Contents"].iter_rows(max_row=5, values_only=True) for v in r if v)
    label = _edition_label(vintage)
    if label not in contents:
        raise ValueError(f"{raw.path.name}: Contents sheet does not mention {label!r}; wrong edition?")
    return wb


def _sheet_topic(ws) -> str:
    """'Table A.5: Current receipts' -> 'current receipts'; '5.1 Debt interest ... ready reckoner' -> 'debt interest ... ready reckoner'."""
    for row in ws.iter_rows(max_row=4, values_only=True):
        for v in row:
            if isinstance(v, str) and v.strip() and "contents" not in v.lower():
                return re.sub(r"^(table\s+)?[a-z]?\.?\d+(\.\d+)*[a-z]?\s*:?\s*", "", v.strip().lower())
    return ""


def _sheet(wb, name: str, topic: str):
    """Sheet `name` if its title is `topic` (regex), else the one sheet whose title is; raises otherwise."""
    if name in wb.sheetnames and re.fullmatch(topic, _sheet_topic(wb[name])):
        return wb[name]
    hits = [ws for ws in wb.worksheets if re.fullmatch(topic, _sheet_topic(ws))]
    if len(hits) != 1:
        raise ValueError(f"obr_efo: expected sheet {name!r} titled {topic!r}; found {[ws.title for ws in hits]}")
    log.warning("obr_efo: sheet %r not found as expected; using %r", name, hits[0].title)
    return hits[0]


def _pick(raws: list[RawArtifact], role: str) -> RawArtifact:
    suffix = FILES[role][1]
    hits = [r for r in raws if r.path.name.endswith(f"_{suffix}")]
    if len(hits) != 1:
        raise ValueError(f"obr_efo: expected one *_{suffix} artifact, got {[r.path.name for r in hits]}")
    return hits[0]


# --------------------------------------------------------------------------- parse


def _obs(t: Table, series_id: str, period: str, value: float, unit: str = "gbp_bn", note: str | None = None) -> Observation:
    return Observation(
        series_id=series_id,
        period=period,
        geography="UK",
        value=round(value, 6),
        unit=unit,
        kind=t.kinds[period],
        source_id=SOURCE.id,
        vintage=t.vintage,
        quality="sourced",
        method_note=note,
    )


def _series(table: Table, row: Row, series_id: str, unit: str = "gbp_bn", note: str | None = None) -> list[Observation]:
    return [_obs(table, series_id, p, row.values[p], unit, note) for p in table.periods.values()]


def _receipts(t: Table) -> list[Observation]:
    """Table A.5 Current receipts."""
    rows = t.data_rows()
    na = t.find(NA_TAXES)
    total = t.find(CURRENT_RECEIPTS)
    label_col = min(c for r in rows for c, _ in r.labels)

    def is_sub(r: Row) -> bool:  # "of which" rows carry their label in a second label column
        return any(c > label_col for c, _ in r.labels)

    top = [r for r in rows if not is_sub(r)]
    tax_rows = [r for r in top if r.index < na.index]
    non_tax_rows = [r for r in top if na.index < r.index < total.index]
    periods = list(t.periods.values())
    out: list[Observation] = []

    # Audit: every OBR row as receipts.detail.<slug>
    parent = None
    for r in rows:
        if r is total:
            continue
        if is_sub(r):
            sid = f"receipts.detail.{slug(parent.label)}_{slug(r.label)}"
            note = f"'Of which' component of the A.5 row '{parent.label}'; not additive with it."
        else:
            parent = r
            sid = f"receipts.detail.{slug(r.label)}"
            note = None
            if r is na:
                note = "A.5 subtotal of all tax rows above it."
            elif r.index > total.index:
                note = "A.5 memo item; not part of current receipts."
        out += [_obs(t, sid, p, r.values[p], note=note) for p in periods if p in r.values]

    # Checks on the published structure before mapping
    for p in periods:
        s_tax = sum(r.values[p] for r in tax_rows)
        if abs(s_tax - na.values[p]) > SUM_TOLERANCE_BN:
            raise ValueError(f"A.5 {p}: tax rows sum {s_tax:.3f} != National Accounts taxes {na.values[p]:.3f}")
        s_all = na.values[p] + sum(r.values[p] for r in non_tax_rows)
        if abs(s_all - total.values[p]) > SUM_TOLERANCE_BN:
            raise ValueError(f"A.5 {p}: taxes + non-tax rows {s_all:.3f} != current receipts {total.values[p]:.3f}")

    out += _series(t, total, "receipts.total")

    by_key: dict[str, Row] = {}
    for r in tax_rows:
        if r.key in by_key:
            raise ValueError(f"A.5: duplicate tax row {r.label!r}")
        by_key[r.key] = r
    used: set[str] = set()
    for line, keys in RECEIPT_LINES.items():
        missing = [k for k in keys if k not in by_key]
        if missing:
            raise ValueError(f"A.5: rows {missing} for receipts.{line} not found; tax rows are {sorted(by_key)}")
        used.update(keys)
        picked = [by_key[k] for k in keys]
        note = None if len(picked) == 1 else "Sum of A.5 rows: " + " + ".join(r.label for r in picked) + "."
        out += [_obs(t, f"receipts.{line}", p, sum(r.values[p] for r in picked), note=note) for p in periods]

    other = [r for r in tax_rows if r.key not in used]
    note = "Sum of the remaining A.5 tax rows: " + ", ".join(r.label for r in other) + "."
    out += [_obs(t, "receipts.other_taxes", p, sum(r.values[p] for r in other), note=note) for p in periods]

    note = "Sum of A.5 non-tax rows: " + ", ".join(r.label for r in non_tax_rows) + "."
    out += [_obs(t, "receipts.non_tax", p, sum(r.values[p] for r in non_tax_rows), note=note) for p in periods]
    return out


def _spending(t: Table) -> list[Observation]:
    """Table A.7 Total managed expenditure."""
    out = _series(t, t.find("Total managed expenditure"), "spending.tme")
    di = t.find("Central government debt interest", startswith=True)
    out += _series(t, di, "spending.debt_interest")
    return out


def _aggregates(t: Table) -> list[Observation]:
    """Table A.9 Fiscal aggregates: a per-cent-of-GDP block, then a £ billion block."""
    out = _series(t, t.find("Public sector net borrowing", section="£ billion"), "fiscal.psnb")
    out += _series(t, t.find("Public sector net debt", section="£ billion"), "fiscal.psnd",
                   note="End-March position.")
    out += _series(t, t.find("Public sector net debt", section="per cent of gdp"), "fiscal.psnd_pct_gdp", unit="pct_gdp",
                   note="End-March position over nominal GDP centred on end-March (A.9 footnote).")
    return out


def _determinants(t: Table) -> list[Observation]:
    """Table A.3 Determinants of the fiscal forecast: nominal GDP."""
    out = _series(t, t.find("Nominal GDP (£ billion)"), "macro.nominal_gdp",
                  note="Fiscal-year nominal GDP, non-seasonally adjusted; OBR's denominator for receipts, spending and deficit ratios.")
    out += _series(t, t.find("Nominal GDP (centred end-March £bn)"), "macro.nominal_gdp_centred_end_march",
                   note="Nominal GDP centred on end-March; OBR's denominator for PSND as a share of GDP.")
    return out


_RR_PP = re.compile(r"^1 percentage point increase in (.+)$", re.I)
_RR_BN = re.compile(r"^£\s*(\d+)\s*bn increase in (.+)$", re.I)


def _ready_reckoner(t: Table) -> list[Observation]:
    """Table 5.1 Debt interest (net of APF) ready reckoner, £bn change by year."""
    out = []
    for r in t.data_rows():
        if m := _RR_PP.match(r.label):
            sid = f"reckoner.debt_interest.{slug(m.group(1))}_1pp"
        elif m := _RR_BN.match(r.label):
            sid = f"reckoner.debt_interest.{slug(m.group(2))}_{m.group(1)}bn"
        else:
            raise ValueError(f"5.1: unrecognised ready-reckoner row {r.label!r}")
        note = f"OBR 5.1: change in central government debt interest net of APF from '{r.label}', assumed from the start of {min(t.periods.values())} and sustained."
        out += [_obs(t, sid, p, r.values[p], note=note) for p in t.periods.values() if p in r.values]
    return out


def parse(raws: list[RawArtifact]) -> list[Observation]:
    annex_raw, rr_raw = _pick(raws, "annex"), _pick(raws, "debt_interest_rr")
    vintages = {_vintage_of(r.path) for r in (annex_raw, rr_raw)} | {r.vintage for r in (annex_raw, rr_raw) if r.vintage}
    if len(vintages) != 1:
        raise ValueError(f"obr_efo: workbooks from different editions {sorted(vintages)}")
    vintage = vintages.pop()
    annex, rr = _workbook(annex_raw, vintage), _workbook(rr_raw, vintage)
    out: list[Observation] = []
    out += _determinants(Table(_sheet(annex, "TA.3", r"determinants of the fiscal forecast"), vintage))
    out += _receipts(Table(_sheet(annex, "TA.5", r"current receipts"), vintage))
    out += _spending(Table(_sheet(annex, "TA.7", r"total managed expenditure"), vintage))
    out += _aggregates(Table(_sheet(annex, "TA.9", r"fiscal aggregates"), vintage))
    out += _ready_reckoner(Table(_sheet(rr, "5.1", r"debt interest\b.*\bready reckoner"), vintage))
    seen = set()
    for o in out:
        k = (o.series_id, o.period)
        if k in seen:
            raise ValueError(f"obr_efo: duplicate observation {k}")
        seen.add(k)
    return out
