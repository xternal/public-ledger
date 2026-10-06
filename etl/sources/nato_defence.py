"""
NATO "Defence Expenditure of NATO Countries" (from 2026 titled "Defence
Investment of NATO Countries (2014-2026)"): the United Kingdom rows.

Why: the sandbox's defence lever is stated as a share of GDP, and politicians
and targets (2%, 2.5%, the 3.5% "core defence" guideline agreed at The Hague
in 2025) use NATO's definition, which counts more than COFOG defence in HMT
PESA (for example pensions paid to retired military and civilian MoD staff).

Emits, geography UK, one observation per year in NATO's table (2014 onwards):

  defence.nato_pct_gdp   Table 3, UK core defence expenditure as a share of GDP (pct_gdp)
  defence.nato_gbp_bn    Table 1, UK core defence expenditure, current prices,
                         million pounds / 1000 (gbp_bn)

Years: NATO's notes (2026 edition, page 14, "Note to readers") say the fiscal
year is designated by the year that holds most of its months, "e.g. 2025
represents the fiscal year 2025/2026 for Canada and the United Kingdom". So
NATO's year Y for the UK is period fiscal_year(Y): 2025 -> "2025-26".

Estimates: years whose column header carries NATO's "e" sign (e.g. "2025e";
the table notes say "Figures for 2025 and 2026 are estimates") are kind
"forecast" with a method_note starting "NATO estimate"; earlier years are
"outturn". All values are read directly from the table (quality "sourced").

Share of GDP: Table 3 is headed "Share of real GDP (%), based on 2021 prices".
NATO deflates defence spending and GDP with the same GDP deflator, so the
ratio equals the current-price share (Table 2 / Table 5, both current US
dollars, give the same figures to rounding; the tests check this). NATO's GDP
comes from DG ECFIN, the IMF and the OECD, not the OBR, so NATO's share times
OBR GDP does not reproduce NATO's pounds exactly.

Finding the edition: the NATO topic page "Defence investment and NATO's 5%
commitment" lists one PDF per edition (".../def-exp-2026-en.pdf"). The Excel
tables sit next to the PDF with the same name and ".xlsx" (the 2025 release
page links ".../def-exp-2025-en.xlsx" as "Download the tables in Excel
format"). The 2026 PDF is a scanned image with no text layer, so only the
Excel file is parsed; if it is missing, fetch() fails rather than guess.

Raw files:
  topic.html              the topic page with the archive of editions
  def-exp-<year>-en.xlsx  the tables of the latest edition (Tables 1, 3, 2 and 5 are read)
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlparse

import httpx
import openpyxl

from etl.core import Observation, RawArtifact, Source, download, fiscal_year, links

SOURCE_ID = "nato_defence"
log = logging.getLogger(__name__)
TOPIC_URL = "https://www.nato.int/en/what-we-do/introduction-to-nato/defence-expenditures-and-natos-5-commitment"
TERMS_URL = "https://www.nato.int/en/about-us/official-texts-and-resources/use-of-nato-content-and-brand"

SOURCE = Source(
    id=SOURCE_ID,
    title="Defence Expenditure of NATO Countries (tables in Excel format)",
    publisher="NATO",
    url=TOPIC_URL,
    # NATO publishes no open licence. Its terms: credit NATO, no sale or advertising use.
    licence=f"NATO terms for external use of content: credit NATO; not to be sold or used for advertising ({TERMS_URL})",
    published_on=date(2026, 7, 9),  # 2026 edition (data cut-off 3 July 2026, Excel saved 9 July); fetch() updates it
    cadence_days=365,  # one edition a year: June 2024, August 2025, July 2026
    grace_days=60,
)

PCT = "defence.nato_pct_gdp"
GBP = "defence.nato_gbp_bn"
COUNTRY = r"United Kingdom"

EDITION_LINK = r"def-exp-(\d{4})(?:-tables)?-en\.(?:pdf|xlsx)$"
YEAR_LABEL = re.compile(r"^(\d{4})\s*(e)?$", re.I)
MISSING = {"..", "-", "–", ""}


# ---------------------------------------------------------------------------
# fetch


def editions(html: str, base: str = TOPIC_URL) -> list[tuple[int, str]]:
    """(edition year, PDF or XLSX url) for every edition linked on the topic page, newest first."""
    out = []
    for url in links(html, base, EDITION_LINK):
        m = re.search(EDITION_LINK, urlparse(url).path, re.I)
        out.append((int(m.group(1)), url))
    return sorted(out, key=lambda t: (t[0], t[1].lower().endswith(".xlsx")), reverse=True)


def xlsx_candidates(year: int, found: list[tuple[int, str]]) -> list[str]:
    """Where the Excel tables of an edition should be: a direct link first, else beside the PDF."""
    out = [u for y, u in found if y == year and u.lower().endswith(".xlsx")]
    for y, u in found:
        if y == year and u.lower().endswith(".pdf"):
            out.append(re.sub(r"\.pdf$", ".xlsx", u, flags=re.I))
            out.append(re.sub(r"-en\.pdf$", "-TABLES-en.xlsx", u, flags=re.I))  # older naming
    return list(dict.fromkeys(out))


def edition_year(raw: RawArtifact) -> int:
    m = re.search(EDITION_LINK, urlparse(raw.url).path, re.I)
    if not m:
        raise ValueError(f"{SOURCE_ID}: cannot read the edition year from {raw.url}")
    return int(m.group(1))


def workbook_date(path: Path, edition: int) -> date | None:
    """When NATO saved the Excel tables (document properties), if it falls in the edition's year."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        for d in (wb.properties.modified, wb.properties.created):
            if isinstance(d, datetime) and d.year == edition:
                return d.date()
    finally:
        wb.close()
    return None


def vintage_for(raw: RawArtifact) -> str:
    edition = edition_year(raw)
    d = workbook_date(raw.path, edition)
    return f"NATO-DE-{edition}-{d.month:02d}" if d else f"NATO-DE-{edition}"


def fetch(since: date | None = None) -> list[RawArtifact]:
    page = download(SOURCE_ID, TOPIC_URL, "topic.html")
    found = editions(page.path.read_text(encoding="utf-8", errors="replace"))
    if not found:
        raise RuntimeError(f"{SOURCE_ID}: no 'def-exp-<year>-en' edition links on {TOPIC_URL}")
    year = found[0][0]
    tried = []
    for url in xlsx_candidates(year, found):
        try:
            raw = download(SOURCE_ID, url, f"def-exp-{year}-en.xlsx")
        except httpx.HTTPStatusError as e:
            if e.response.status_code in (403, 404, 410):
                tried.append(f"{url} ({e.response.status_code})")
                continue
            raise
        raw.vintage = vintage_for(raw)
        published = workbook_date(raw.path, year)
        SOURCE.title = f"Defence Expenditure of NATO Countries, {year} edition (tables in Excel format)"
        if published:
            SOURCE.published_on = published
        log.info("nato_defence: using the %s edition (%s, saved %s) from %s", year, raw.vintage, published, url)
        return [page, raw]
    raise RuntimeError(
        f"{SOURCE_ID}: the {year} edition is listed on the topic page but no Excel tables were found "
        f"(tried {', '.join(tried) or 'nothing'}). The PDF is not parsed: the 2026 one is a scanned image."
    )


# ---------------------------------------------------------------------------
# parse


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def table_sheet(wb, number: int):
    """The sheet whose first cell reads 'Table <number>: ...'."""
    found = [ws for ws in wb.worksheets if re.match(rf"^Table {number}\s*:", _text(ws.cell(1, 1).value), re.I)]
    if len(found) != 1:
        raise ValueError(f"{SOURCE_ID}: expected one sheet titled 'Table {number}: ...', found {len(found)}")
    return found[0]


def block(ws, heading: str) -> tuple[list[tuple[int, bool]], dict[str, list]]:
    """
    One block of a NATO table: the first row whose column A matches `heading`,
    then the year header row (column A empty, '2014' ... '2026e'), then one row
    per country until a blank row. Returns ([(year, is_estimate)], {label: values}).
    """
    rows = list(ws.iter_rows(values_only=True))
    start = next((i for i, r in enumerate(rows) if re.search(heading, _text(r[0]), re.I)), None)
    if start is None:
        raise ValueError(f"{SOURCE_ID}: {ws.title}: no row matching {heading!r} in column A")
    hdr = next((i for i in range(start + 1, min(start + 4, len(rows))) if not _text(rows[i][0]) and _text(rows[i][1] if len(rows[i]) > 1 else None)), None)
    if hdr is None:
        raise ValueError(f"{SOURCE_ID}: {ws.title}: no year header under {heading!r}")
    years: list[tuple[int, bool]] = []
    for cell in rows[hdr][1:]:
        if not _text(cell):
            break
        m = YEAR_LABEL.match(_text(cell))
        if not m:
            raise ValueError(f"{SOURCE_ID}: {ws.title}: unexpected year label {cell!r}")
        years.append((int(m.group(1)), bool(m.group(2))))
    ys = [y for y, _ in years]
    if len(ys) < 5 or ys != list(range(ys[0], ys[0] + len(ys))):
        raise ValueError(f"{SOURCE_ID}: {ws.title}: years are not consecutive: {ys}")
    flags = [e for _, e in years]
    if flags != sorted(flags):
        raise ValueError(f"{SOURCE_ID}: {ws.title}: estimate years are not the latest ones: {years}")
    body: dict[str, list] = {}
    for r in rows[hdr + 1:]:
        label = _text(r[0])
        if not label:
            break
        body[label] = list(r[1:1 + len(years)])
    if len(body) < 10:
        raise ValueError(f"{SOURCE_ID}: {ws.title}: only {len(body)} rows under {heading!r}")
    return years, body


def country_row(body: dict[str, list], pattern: str, where: str) -> list:
    hits = [k for k in body if re.match(pattern, k)]
    if len(hits) != 1:
        raise ValueError(f"{SOURCE_ID}: {where}: expected one row matching {pattern!r}, found {hits}")
    return body[hits[0]]


def number(v, where: str) -> float | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    t = _text(v).replace(",", "")
    if t in MISSING:
        return None  # NATO's ".." (not available) or "-" (nil)
    try:
        return float(t)
    except ValueError:
        raise ValueError(f"{SOURCE_ID}: {where}: not a number: {v!r}") from None


def estimate_note(wb) -> set[int] | None:
    """Years named in the notes' 'Figures for 2025 and 2026 are estimates', if present."""
    for ws in wb.worksheets:
        for (v,) in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            m = re.search(r"Figures for ([\d\s,and]+?) are estimates", _text(v), re.I)
            if m:
                return {int(y) for y in re.findall(r"\d{4}", m.group(1))}
    return None


def uk_tables(path: Path) -> dict:
    """The UK rows of Tables 1, 2, 3 and 5, with the year header and table titles."""
    wb = openpyxl.load_workbook(path, data_only=True)
    t1, t2, t3, t5 = (table_sheet(wb, n) for n in (1, 2, 3, 5))
    for ws, unit in ((t1, r"million national currency"), (t2, r"million us dollars"), (t5, r"million us dollars")):
        if not re.search(unit, _text(ws.cell(2, 1).value), re.I):
            raise ValueError(f"{SOURCE_ID}: {ws.title}: unit row reads {ws.cell(2, 1).value!r}, expected {unit!r}")
    if not re.search(r"defence expenditure", _text(t1.cell(1, 1).value), re.I) or not re.search(r"share of GDP", _text(t3.cell(1, 1).value), re.I):
        raise ValueError(f"{SOURCE_ID}: unexpected table titles {t1.cell(1, 1).value!r} / {t3.cell(1, 1).value!r}")

    y1, b1 = block(t1, r"^current prices$")
    y2, b2 = block(t2, r"^current prices and exchange rates$")
    y3, b3 = block(t3, r"^share of (real )?GDP")
    y5, b5 = block(t5, r"^current prices and exchange rates$")
    if not (y1 == y2 == y3 == y5):
        raise ValueError(f"{SOURCE_ID}: Tables 1, 2, 3 and 5 have different years: {y1} / {y2} / {y3} / {y5}")
    note = estimate_note(wb)
    flagged = {y for y, e in y1 if e}
    if note is not None and note != flagged:
        raise ValueError(f"{SOURCE_ID}: notes say {sorted(note)} are estimates but the headers flag {sorted(flagged)}")
    return {
        "years": y1,
        # "Table 3: Core defence expenditure as a share of GDP and annual real change" -> without the second half.
        "titles": {n: re.sub(r"\s+and annual real change$", "", _text(ws.cell(1, 1).value), flags=re.I) for n, ws in ((1, t1), (3, t3))},
        "gbp_m": country_row(b1, rf"^{COUNTRY}\b.*\(Pounds\)$", "Table 1"),
        "usd_m": country_row(b2, rf"^{COUNTRY}\b", "Table 2"),
        "pct_gdp": country_row(b3, rf"^{COUNTRY}\b", "Table 3"),
        "gdp_usd_m": country_row(b5, rf"^{COUNTRY}\b", "Table 5"),
    }


def parse(raws: list[RawArtifact]) -> list[Observation]:
    xlsx = [r for r in raws if r.path.suffix.lower() == ".xlsx"]
    if len(xlsx) != 1:
        raise ValueError(f"{SOURCE_ID}: expected one Excel file among the raw files, got {len(xlsx)}")
    raw = xlsx[0]
    edition = edition_year(raw)
    vintage = raw.vintage or vintage_for(raw)
    t = uk_tables(raw.path)
    years = t["years"]
    if years[-1][0] != edition:
        raise ValueError(f"{SOURCE_ID}: the {edition} edition's tables end in {years[-1][0]}")

    out: list[Observation] = []
    for i, (y, est) in enumerate(years):
        fy = fiscal_year(y)
        label = f"{y}e" if est else str(y)
        mapping = f"NATO's {y} is the UK fiscal year {fy}."
        prefix = f"NATO estimate ('{label}' in the {edition} edition). " if est else ""
        pct = number(t["pct_gdp"][i], f"Table 3 {label}")
        gbp_m = number(t["gbp_m"][i], f"Table 1 {label}")
        if pct is not None:
            if not 0.5 <= pct <= 10:
                raise ValueError(f"{SOURCE_ID}: implausible UK share of GDP {pct} in {label}")
            out.append(Observation(
                series_id=PCT, period=fy, geography="UK", value=round(pct, 3), unit="pct_gdp",
                kind="forecast" if est else "outturn", source_id=SOURCE_ID, vintage=vintage, quality="sourced",
                method_note=f"{prefix}{t['titles'][3]}, NATO definition. {mapping}",
            ))
        if gbp_m is not None:
            if not 10_000 <= gbp_m <= 500_000:
                raise ValueError(f"{SOURCE_ID}: implausible UK spending £{gbp_m}m in {label}")
            out.append(Observation(
                series_id=GBP, period=fy, geography="UK", value=round(gbp_m / 1000, 3), unit="gbp_bn",
                kind="forecast" if est else "outturn", source_id=SOURCE_ID, vintage=vintage, quality="sourced",
                method_note=f"{prefix}{t['titles'][1]}, current prices, £ million / 1000, NATO definition. {mapping}",
            ))
    if not any(o.series_id == PCT for o in out):
        raise ValueError(f"{SOURCE_ID}: no UK share of GDP read")
    out.sort(key=lambda o: (o.series_id, o.period))
    return out
