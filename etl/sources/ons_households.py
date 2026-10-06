"""
ONS Families and households in the UK: number of households, UK, by year.

Two requests per run: the dataset page (the spreadsheet's file name changes
with each edition, e.g. familiesandhouseholdsuk2025.xlsx) and the spreadsheet.

The figure is "All households" in Table 7 (households by type; Table 5 by size
has the same total), in thousands, from the Labour Force Survey household
dataset for April to June of each year. It is a survey estimate: the 95%
confidence interval (about +/-0.15m) is kept in method_note. The official
household estimates for Scotland, Wales and Northern Ireland are published by
NRS, the Welsh Government and NISRA; ONS recommends this table for UK totals.
"""

from __future__ import annotations

import re
from datetime import date, datetime

import openpyxl

from etl.core import Observation, RawArtifact, Source, download, links

SOURCE = Source(
    id="ons_households",
    title="ONS Families and households in the UK",
    publisher="ONS",
    url="https://www.ons.gov.uk/peoplepopulationandcommunity/birthsdeathsandmarriages/families/datasets/familiesandhouseholdsfamiliesandhouseholds",
    licence="OGL v3",
    cadence_days=365,
    grace_days=90,
)

SERIES_ID = "people.households"


def fetch(since: date | None = None) -> list[RawArtifact]:
    """Find the current edition's xlsx on the dataset page and download it (each cached 20 h)."""
    page = download(SOURCE.id, SOURCE.url, "dataset_page.html")
    html = page.path.read_text(errors="replace")
    found = links(html, SOURCE.url, r"file\?uri=.*familiesandhouseholdsfamiliesandhouseholds/current/.*\.xlsx$")
    if not found:
        raise ValueError("ONS families and households: no current .xlsx link on the dataset page")
    url = found[0]
    name = re.sub(r"[^A-Za-z0-9._-]", "_", url.rsplit("/", 1)[-1])
    raw = download(SOURCE.id, url, name)
    raw.vintage = vintage_for(published(raw))
    return [raw]


def published(raw: RawArtifact) -> date:
    """'Date published: 17 April 2026.' on the cover sheet."""
    wb = openpyxl.load_workbook(raw.path, read_only=True, data_only=True)
    try:
        for row in wb.worksheets[0].iter_rows(values_only=True, max_row=15):
            for c in row:
                m = re.search(r"Date published:\s*(\d{1,2} \w+ \d{4})", str(c or ""))
                if m:
                    return datetime.strptime(m.group(1), "%d %B %Y").date()
    finally:
        wb.close()
    raise ValueError("ONS families and households: no 'Date published' on the cover sheet")


def vintage_for(d: date) -> str:
    return f"FH-{d:%Y-%m}"


def read_households(path) -> dict[int, tuple[float, float | None]]:
    """{year: (households in thousands, 95% CI +/- in thousands)} from the 'Households by type' table."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = sorted(
            wb.worksheets,
            key=lambda ws: 0 if str(next(ws.iter_rows(values_only=True, max_row=1))[0] or "").startswith("Table 7") else 1,
        )
        for ws in sheets:
            title = str(next(ws.iter_rows(values_only=True, max_row=1))[0] or "")
            if not re.match(r"Table \d+: Households by", title):
                continue
            header: list | None = None
            for row in ws.iter_rows(values_only=True):
                first = str(row[0] or "").strip()
                if first.startswith("Number of households (thousands)"):
                    header = list(row)
                elif header and first == "All households":
                    out: dict[int, tuple[float, float | None]] = {}
                    for i, h in enumerate(header):
                        m = re.fullmatch(r"(\d{4}) Estimate", str(h or "").strip())
                        if m and isinstance(row[i], (int, float)):
                            ci = row[i + 2] if i + 2 < len(row) and str(header[i + 2] or "").endswith("CI+/-") else None
                            out[int(m.group(1))] = (float(row[i]), float(ci) if isinstance(ci, (int, float)) else None)
                    if out:
                        return out
    finally:
        wb.close()
    raise ValueError("ONS families and households: no 'All households' row found")


def parse(raws: list[RawArtifact]) -> list[Observation]:
    raw = raws[0] if isinstance(raws, list) else raws
    pub = published(raw)
    SOURCE.published_on = pub
    vintage = raw.vintage or vintage_for(pub)
    out = []
    for year, (thousands, ci) in sorted(read_households(raw.path).items()):
        note = f"All households, UK, Labour Force Survey April-June {year}; thousands / 1000"
        note += f"; 95% confidence interval +/-{ci / 1000:.3f}m." if ci is not None else "."
        out.append(Observation(
            series_id=SERIES_ID, period=str(year), value=thousands / 1000.0, unit="households_m", kind="outturn",
            source_id=SOURCE.id, vintage=vintage, quality="sourced", method_note=note,
        ))
    return out
