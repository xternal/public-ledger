"""
ONS mid-year population estimates: births and deaths in the UK, year to each mid-year.

Discovery (fetch): the dataset page "Population estimates for the UK, England, Wales,
Scotland and Northern Ireland" links one "Mid-2011 to mid-<year>" detailed time
series workbook per edition; the newest is used. Its sheet MYEB5 holds the
components of population change for the UK and its countries, one column per
component and year (births_2012, deaths_2012, ...). We read the UNITED KINGDOM
row (code K02000001), by column name, never by position.

These are the past values shown next to the ONS national population projections,
which count births and deaths on the same basis (the year to 30 June). The
workbook is large (about 50 MB); core.download() caches it for 20 hours.
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

from etl.core import Observation, RawArtifact, Source, download, links

log = logging.getLogger(__name__)

DATASET_PAGE = (
    "https://www.ons.gov.uk/peoplepopulationandcommunity/populationandmigration/populationestimates/"
    "datasets/populationestimatesforukenglandandwalesscotlandandnorthernireland"
)
UK_CODE = "K02000001"
SHEET = "MYEB5"
COMPONENTS = {"births": "people.births", "deaths": "people.deaths"}

SOURCE = Source(
    id="ons_mye_components",
    title="ONS Population estimates for the UK: components of population change (table MYEB5)",
    publisher="ONS",
    url=DATASET_PAGE,
    licence="OGL v3",
    published_on=date(2026, 10, 1),
    cadence_days=365,
    grace_days=90,
)


def vintage_for(d: date) -> str:
    return f"MYEB-{d:%Y-%m-%d}"


def latest_time_series(html: str, base: str = DATASET_PAGE) -> tuple[int, str]:
    """(last mid-year, URL) of the newest 'mid-2011 to mid-<year>' workbook linked on the dataset page."""
    found = []
    for u in links(html, base, r"/file\?uri=.*/mid2011tomid\d{4}/[^/]+\.xlsx$"):
        m = re.search(r"/mid2011tomid(\d{4})/", u)
        if m:
            found.append((int(m.group(1)), u))
    if not found:
        raise LookupError(f"no mid-2011 to mid-year time series workbook on {base}")
    return max(found)


def fetch(since: date | None = None) -> list[RawArtifact]:
    html = download(SOURCE.id, DATASET_PAGE, "dataset.html").path.read_text(errors="replace")
    year, url = latest_time_series(html)
    raw = download(SOURCE.id, url, f"myeb_2011_{year}.xlsx")
    raw.vintage = vintage_for(read(raw.path)[0])
    return [raw]


def read(path: Path) -> tuple[date, dict[str, dict[int, float]]]:
    """(publication date, {component: {year: persons}}) for the UK from a detailed time series workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cover = " ".join(str(v) for row in wb["Cover sheet"].iter_rows(values_only=True) for v in row if v)
    m = re.search(r"Date published:\s*(\d{1,2} \w+ \d{4})", cover)
    if not m:
        raise ValueError(f"{path.name}: no publication date on the cover sheet")
    published = datetime.strptime(m.group(1), "%d %B %Y").date()
    if SHEET not in wb.sheetnames:
        raise ValueError(f"{path.name}: no sheet {SHEET}")
    header: list[str] | None = None
    values: dict[str, dict[int, float]] = {c: {} for c in COMPONENTS}
    for row in wb[SHEET].iter_rows(values_only=True):
        if row and row[0] == "Code":
            header = [str(v) if v is not None else "" for v in row]
            continue
        if header and row and row[0] == UK_CODE:
            for name, v in zip(header, row):
                m = re.fullmatch(r"(births|deaths)_(\d{4})", name)
                if m and isinstance(v, (int, float)):
                    values[m.group(1)][int(m.group(2))] = float(v)
            break
    if header is None:
        raise ValueError(f"{path.name}: {SHEET} has no header row starting 'Code'")
    for c, by_year in values.items():
        if not by_year:
            raise ValueError(f"{path.name}: no UK {c} in {SHEET}")
    return published, values


def parse(raws: list[RawArtifact]) -> list[Observation]:
    raw = raws[0] if isinstance(raws, list) else raws
    published, values = read(raw.path)
    SOURCE.published_on = published
    vintage = raw.vintage or vintage_for(published)
    out = []
    for comp, series_id in COMPONENTS.items():
        for year, persons in sorted(values[comp].items()):
            out.append(Observation(
                series_id=series_id, period=str(year), value=persons / 1000, unit="persons_k", kind="outturn",
                source_id=SOURCE.id, vintage=vintage, quality="sourced",
                method_note=f"UK {comp}, year to mid-{year} (1 July {year - 1} to 30 June {year}); ONS table {SHEET}, persons / 1,000.",
            ))
    return out
