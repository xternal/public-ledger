"""
ONS national population projections (NPP), UK: the principal projection and every published variant.

Discovery (fetch):
  1. The NPP table-of-contents dataset page lists one workbook per edition
     ("2024-based national population projections table of contents"). The newest
     base year is used.
  2. That workbook's "Links" sheet names, for every variant code (ppp, hpp, lpp, ...),
     the dataset page of the UK "components of change summary table".
  3. Each dataset page links one workbook per edition; the one under "/<base>based/"
     is downloaded. A variant whose page has no file for this edition is left out
     (logged), never filled in; the principal projection must be there.

Variant codes are ONS's own: three letters for fertility, life expectancy (mortality)
and net migration, each h(igh), l(ow) or p(rincipal); z = zero net migration,
r = replacement fertility, n = no long-term mortality improvement.

Each summary workbook (sheet PERSONS) has three tables read by row label and by the
year in the header, never by cell position:

  Table 1  births, deaths, net migration (thousands, year to mid-year), total
           fertility rate, life expectancy at birth by sex
  Table 2  population at mid-year by age group (we keep "All ages")
  Table 3  working age and pension age (thousands, at mid-year, using the
           state pension age in law for that year) and pension age per 1,000 of
           working age (ONS's old-age dependency ratio)

Years run from the base year (a mid-year estimate, kind "outturn") to LAST_YEAR,
the horizon of the OBR's long-term projections; ONS projects 100 years ahead.
"""

from __future__ import annotations

import logging
import re
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

import openpyxl

import httpx

from etl.core import Observation, RawArtifact, Source, download, links

log = logging.getLogger(__name__)

TOC_PAGE = (
    "https://www.ons.gov.uk/peoplepopulationandcommunity/populationandmigration/populationprojections/"
    "datasets/2014basednationalpopulationprojectionstableofcontents"
)
LAST_YEAR = 2075

SOURCE = Source(
    id="ons_npp",
    title="ONS National population projections: 2024-based (UK principal and variant projections)",
    publisher="ONS",
    url="https://www.ons.gov.uk/peoplepopulationandcommunity/populationandmigration/populationprojections/bulletins/nationalpopulationprojections/2024based",
    licence="OGL v3",
    published_on=date(2026, 4, 28),
    # A full set every two years or so (2022-based in January 2025, 2024-based in April 2026).
    cadence_days=900,
    grace_days=180,
)

BULLETIN = "https://www.ons.gov.uk/peoplepopulationandcommunity/populationandmigration/populationprojections/bulletins/nationalpopulationprojections/{base}based"

# What each letter of a variant code means, by position: fertility, life expectancy, migration.
ASSUMPTION_LETTERS = (
    {"p": "principal", "h": "high", "l": "low", "r": "replacement"},
    {"p": "principal", "h": "high", "l": "low", "n": "no_improvement"},
    {"p": "principal", "h": "high", "l": "low", "z": "zero"},
)
CODE = re.compile(r"^[phlr][phln][phlz]$")

# measure -> (table number, exact row label, unit)
ROWS: dict[str, tuple[int, str, str]] = {
    "births": (1, "births", "persons_k"),
    "deaths": (1, "deaths", "persons_k"),
    "net_migration": (1, "net migration", "persons_k"),
    "tfr": (1, "total fertility rate (tfr)", "children_per_woman"),
    "life_expectancy_male": (1, "eolb males", "years"),
    "life_expectancy_female": (1, "eolb females", "years"),
    "population": (2, "all ages", "persons_k"),
    "working_age": (3, "working age", "persons_k"),
    "pension_age": (3, "pension age", "persons_k"),
    "oadr": (3, "pension age per 1,000 persons of working age", "per_1000"),
}
# Rows whose base-year value carries ONS's note on state pension age.
SPA_MEASURES = {"working_age", "pension_age", "oadr"}

_FILE = re.compile(r"^npp(\d{4})_([a-z]{3})_uk_summary\.xlsx$")


# ONS answers 429 to bursts; wait and ask again, at most this many times per file.
RETRIES = 4


def _download(url: str, filename: str, **kw) -> RawArtifact:
    for attempt in range(RETRIES + 1):
        try:
            return download(SOURCE.id, url, filename, **kw)
        except httpx.HTTPStatusError as e:
            if e.response.status_code != 429 or attempt == RETRIES:
                raise
            wait = float(e.response.headers.get("retry-after") or 0) or 5.0 * (attempt + 1)
            log.info("ons_npp: ONS asked us to slow down; waiting %.0fs before %s", wait, filename)
            time.sleep(min(wait, 60.0))
    raise AssertionError("unreachable")


def vintage_for(base: int, published: date) -> str:
    return f"NPP-{base}based-{published:%Y-%m-%d}"


# --------------------------------------------------------------------------- discovery


def toc_editions(html: str, base_url: str = TOC_PAGE) -> list[tuple[int, str]]:
    """(base year, workbook URL) for every edition linked on the table-of-contents page, newest first; full editions before interim ones."""
    found = []
    for u in links(html, base_url, r"/file\?uri=.*nationalpopulationprojectionstableofcontents.*\.xlsx?$"):
        # The edition is the folder the file sits in; the dataset's own slug starts "2014based" for every edition.
        folder = u.rsplit("/", 2)[-2]
        m = re.match(r"(\d{4})based(interim)?nationalpopulationprojectionstableofcontents", folder)
        if m:
            found.append((int(m.group(1)), not m.group(2), u))
    found.sort(reverse=True)
    return [(b, u) for b, _, u in found]


def toc_summary_pages(path: Path) -> tuple[date, dict[str, str]]:
    """(publication date, {variant code: UK summary dataset page}) from a table-of-contents workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    published = None
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and (m := re.search(r"Date of publication:\s*(\d{1,2} \w+ \d{4})", v)) and published is None and ws.title != "Earlier NPP releases":
                    published = datetime.strptime(m.group(1), "%d %B %Y").date()
    if "Links" not in wb.sheetnames:
        raise LookupError(f"{path.name}: no Links sheet")
    codes: list[str] | None = None
    pages: dict[str, str] = {}
    for row in wb["Links"].iter_rows(values_only=True):
        vals = [v for v in row if v not in (None, "")]
        if vals and all(isinstance(v, str) and CODE.match(v) for v in vals) and "ppp" in vals:
            codes = [str(v) for v in vals]
        elif codes and len(vals) >= 2 and vals[0] == "UK" and "summary" in str(vals[1]).lower():
            pages = {c: str(u) for c, u in zip(codes, vals[2:]) if str(u).startswith("https://www.ons.gov.uk/")}
    if published is None:
        raise LookupError(f"{path.name}: no publication date")
    if "ppp" not in pages:
        raise LookupError(f"{path.name}: no UK summary table for the principal projection")
    return published, pages


def edition_file(html: str, page_url: str, base: int) -> str | None:
    """The workbook of this base year linked on a summary dataset page, or None."""
    found = links(html, page_url, rf"/file\?uri=.*/{base}based/[^/]+\.xlsx?$")
    return found[0] if found else None


def fetch(since: date | None = None) -> list[RawArtifact]:
    toc_html = _download(TOC_PAGE, "toc.html").path.read_text(errors="replace")
    editions = toc_editions(toc_html)
    if not editions:
        raise LookupError(f"no NPP edition on {TOC_PAGE}")
    base, toc_url = editions[0]
    toc = _download(toc_url, f"npp{base}_toc.xlsx")
    published, pages = toc_summary_pages(toc.path)
    vintage = vintage_for(base, published)
    SOURCE.title = f"ONS National population projections: {base}-based (UK principal and variant projections)"
    SOURCE.url = BULLETIN.format(base=base)
    SOURCE.published_on = published
    log.info("ons_npp: %s-based projections published %s, %d variants listed", base, published, len(pages))
    raws = []
    for code, page in pages.items():
        html = _download(page, f"npp_{code}_uk_summary.html").path.read_text(errors="replace")
        url = edition_file(html, page, base)
        if url is None:
            if code == "ppp":
                raise LookupError(f"{page}: no {base}-based workbook for the principal projection")
            log.warning("ons_npp: %s has no %s-based workbook; variant %s left out", page, base, code)
            continue
        raws.append(_download(url, f"npp{base}_{code}_uk_summary.xlsx", vintage=vintage))
    return raws


# --------------------------------------------------------------------------- parsing


def _clean(v) -> str:
    return re.sub(r"\s*\[note \d+\]", "", str(v)).strip().lower() if v is not None else ""


def _year(v) -> int | None:
    m = re.match(r"^\s*(\d{4})\b", str(v)) if v is not None else None
    return int(m.group(1)) if m else None


def read_summary(path: Path) -> dict:
    """{'base', 'published', 'projection', 'tables': {n: {label: {year: value}}}, 'notes': {n: text}} from a UK summary workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cover = " ".join(str(v) for row in wb["Cover sheet"].iter_rows(values_only=True) for v in row if v)
    m_base = re.search(r"National population projections:\s*(\d{4})-based", cover)
    m_date = re.search(r"Date published:\s*(\d{1,2} \w+ \d{4})", cover)
    m_proj = re.search(r"Projection:\s*([^\n]+?)(?:\s+This spreadsheet|$)", cover)
    m_cov = re.search(r"Coverage:\s*United Kingdom", cover)
    if not (m_base and m_date and m_cov):
        raise ValueError(f"{path.name}: cover sheet does not describe a UK national population projection")
    tables: dict[int, dict[str, dict[int, float]]] = {}
    current: int | None = None
    years: list[int | None] = []
    for row in wb["PERSONS"].iter_rows(values_only=True):
        first = row[0] if row else None
        if isinstance(first, str) and (m := re.match(r"Table (\d+):", first.strip())):
            current, years = int(m.group(1)), []
            tables[current] = {}
            continue
        if current is None or first is None:
            continue
        ys = [_year(v) for v in row[1:]]
        if not years and sum(y is not None for y in ys) >= 3:
            years = ys
            continue
        label = _clean(first)
        if years and label and label not in tables[current]:
            tables[current][label] = {y: float(v) for y, v in zip(years, row[1:]) if y is not None and isinstance(v, (int, float))}
    notes = {}
    if "Notes" in wb.sheetnames:
        for row in wb["Notes"].iter_rows(values_only=True):
            if row and isinstance(row[0], str) and (m := re.match(r"note (\d+)$", row[0].strip())) and len(row) > 1 and row[1]:
                notes[int(m.group(1))] = str(row[1]).strip()
    return {
        "base": int(m_base.group(1)),
        "published": datetime.strptime(m_date.group(1), "%d %B %Y").date(),
        "projection": m_proj.group(1).strip() if m_proj else "",
        "tables": tables,
        "notes": notes,
    }


def code_of(path: Path) -> str:
    m = _FILE.match(path.name)
    if not m:
        raise ValueError(f"unexpected file name {path.name}; expected npp<base>_<code>_uk_summary.xlsx")
    return m.group(2)


def spa_note(notes: dict[int, str]) -> str | None:
    """ONS's note on how working and pension age follow the state pension age."""
    return next((t for t in notes.values() if "state pension age" in t.lower()), None)


def parse(raws: list[RawArtifact]) -> list[Observation]:
    raws = raws if isinstance(raws, list) else [raws]
    out: list[Observation] = []
    codes = []
    for raw in raws:
        code = code_of(raw.path)
        s = read_summary(raw.path)
        vintage = raw.vintage or vintage_for(s["base"], s["published"])
        SOURCE.published_on = s["published"]
        spa = spa_note(s["notes"])
        codes.append(code)
        for measure, (table, label, unit) in ROWS.items():
            series = s["tables"].get(table, {}).get(label)
            if not series:
                raise ValueError(f"{raw.path.name}: no row {label!r} in table {table}")
            for year, value in sorted(series.items()):
                if year > LAST_YEAR:
                    continue
                base_year = year == s["base"]
                note = None
                if base_year:
                    note = f"Base year of the {s['base']}-based projections: ONS mid-{year} population estimate."
                    if measure in SPA_MEASURES and spa:
                        note += f" ONS: {spa}"
                out.append(Observation(
                    series_id=f"people.npp.{code}.{measure}", period=str(year), value=value, unit=unit,
                    kind="outturn" if base_year else "projection", source_id=SOURCE.id, vintage=vintage,
                    quality="sourced", method_note=note,
                ))
    if "ppp" not in codes:
        raise ValueError("no principal projection (ppp) among the NPP workbooks")
    log.info("ons_npp: %d observations for variants %s", len(out), ", ".join(sorted(codes)))
    return out
